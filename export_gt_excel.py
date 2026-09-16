# -*- coding: utf-8 -*-
"""
将 CloudCompare 实测的圆心真值（gt_manual.json）与法向量（cc_norm_manual.json）
汇总导出为 Excel：测量数据汇总.xlsx
用法：python export_gt_excel.py
"""
import json, os, math
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
GT = json.load(open(os.path.join(HERE, "gt_manual.json"), encoding="utf-8"))
CC = json.load(open(os.path.join(HERE, "cc_norm_manual.json"), encoding="utf-8"))
OUT = os.path.join(HERE, "测量数据汇总.xlsx")


def pitch_yaw(n):
    n = list(map(float, n))
    if n[2] < 0:
        n = [-v for v in n]
    return (math.degrees(math.atan2(n[1], n[2])),
            math.degrees(math.atan2(n[0], n[2])))


scenes = sorted((k for k in GT if not k.startswith("_")),
                key=lambda s: int(s.replace("ICP_input", "")))

wb = Workbook()
FILL = PatternFill("solid", fgColor="D9E1F2")


def style_header(ws):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = FILL
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")


ws1 = wb.active
ws1.title = "圆心真值(CC手选)"
ws1.append(["场景", "ly_X", "ly_Y", "ly_Z", "ry_X", "ry_Y", "ry_Z"])
for s in scenes:
    ws1.append([s] + GT[s]["ly"] + GT[s]["ry"])
style_header(ws1)
for row in ws1.iter_rows(min_row=2, min_col=2, max_col=7):
    for cell in row:
        cell.number_format = "0.0000"
ws1.column_dimensions["A"].width = 14
for col in "BCDEFG":
    ws1.column_dimensions[col].width = 12
ws1.freeze_panes = "A2"

ws2 = wb.create_sheet("法向量实测(CC拟合平面)")
ws2.append(["场景", "ly_nx", "ly_ny", "ly_nz", "ly_pitch°", "ly_yaw°",
            "ry_nx", "ry_ny", "ry_nz", "ry_pitch°", "ry_yaw°"])
for s in scenes:
    nly, nry = CC[s]["ly"]["n"], CC[s]["ry"]["n"]
    ply, yly = pitch_yaw(nly)
    pry, yry = pitch_yaw(nry)
    ws2.append([s] + nly + [round(ply, 2), round(yly, 2)]
               + nry + [round(pry, 2), round(yry, 2)])
style_header(ws2)
ang_cols = {5, 6, 10, 11}
for row in ws2.iter_rows(min_row=2, min_col=2, max_col=11):
    for cell in row:
        cell.number_format = "0.00" if cell.column in ang_cols else "0.0000000"
ws2.column_dimensions["A"].width = 14
for i in range(2, 12):
    ws2.column_dimensions[get_column_letter(i)].width = 11
ws2.freeze_panes = "A2"

ws3 = wb.create_sheet("说明")
for text in [
    "测量数据汇总说明",
    "",
    "1. 圆心真值来自 gt_manual.json：用户在 CloudCompare 中手选（input1/2/3/4 的 ry 于 2026-09-15 按 val2 方法重选）。",
    "2. 法向量来自 cc_norm_manual.json：用户在 CloudCompare Tools>Fit>Plane 拟合平面量取（2026-09-16）。",
    "   其中 input0-ry / input8-ly / input9-ly 为纯盘面点复测，input0-ly / input7-ly 为二次复测。",
    "3. 端口定义：ly = 回收口（小圆，半径约 59mm）；ry = 加注口（大圆，半径约 79mm）。",
    "4. 角度约定（与 val2.png 一致，法向统一朝 +z）：pitch = atan2(ny, nz)，yaw = atan2(nx, nz)，单位为度。",
    "5. 坐标为各帧 ICP 配准后点云坐标系，单位 mm；法向量为单位向量。",
]:
    ws3.append([text])
ws3["A1"].font = Font(bold=True, size=13)
for r in range(3, 10):
    ws3.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
ws3.column_dimensions["A"].width = 110

wb.save(OUT)
print("已生成:", OUT)
